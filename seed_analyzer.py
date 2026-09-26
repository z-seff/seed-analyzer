#!/usr/bin/env python3
"""
Seed Analyzer — derives addresses from a BIP39 seed phrase and checks balances.

Schemes (defaults match https://iancoleman.io/bip39/):
  EVM  BIP44   m/44'/60'/0'/0/i    one address is checked across all EVM networks
  TRX  BIP44   m/44'/195'/0'/0/i
  BTC  BIP44   m/44'/0'/0'/0/i     P2PKH   (1...)
  BTC  BIP49   m/49'/0'/0'/0/i     P2SH    (3...)
  BTC  BIP84   m/84'/0'/0'/0/i     P2WPKH  (bc1...)
  BTC  BIP32   m/0/i               P2PKH   (1...)    — BIP32 tab, default path m/0
  BTC  BIP141  m/0/i               P2SH    (3...)    — BIP141 tab, semantics "P2WPKH nested in P2SH"

EVM networks: Ethereum, Arbitrum, Base, Polygon, Avalanche, OP Mainnet —
native coin + USDC/USDT (+ DAI / WETH / bridged .e versions where available).

Output — xlsx: "Addresses" sheet (empty addresses hidden) and "Summary" sheet.

python3.12 -m venv .venv && .venv/bin/pip install -r requirements.txt
.venv/bin/python seed_analyzer.py          # phrase entered via a hidden prompt
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from dataclasses import dataclass, field
from decimal import Decimal
from functools import partial
from typing import Callable

import requests
from bip_utils import (
    Bip32Slip10Secp256k1,
    Bip39MnemonicValidator,
    Bip39SeedGenerator,
    Bip44,
    Bip44Changes,
    Bip44Coins,
    Bip49,
    Bip49Coins,
    Bip84,
    Bip84Coins,
    CoinsConf,
    P2PKHAddrEncoder,
    P2SHAddrEncoder,
    WifEncoder,
)

# ---------------------------------------------------------------- configuration

TRON_API = "https://api.trongrid.io"
BTC_API = "https://blockstream.info/api"

HTTP_TIMEOUT = 30
HTTP_RETRIES = 3
USER_AGENT = "seed-analyzer/1.1"


@dataclass(frozen=True)
class EvmChain:
    key: str
    name: str
    chain_id: int
    rpc: str
    symbol: str
    tokens: dict[str, tuple[str, int]]  # symbol -> (contract, decimals)


# The EVM address is the same across all networks; only the RPC and token contracts differ.
EVM_CHAINS: dict[str, EvmChain] = {
    "eth": EvmChain(
        "eth", "Ethereum", 1, "https://ethereum-rpc.publicnode.com", "ETH",
        {
            "USDT": ("0xdAC17F958D2ee523a2206206994597C13D831ec7", 6),
            "USDC": ("0xA0b86991c6218b36c1d19D4a2e9Eb0cE3606eB48", 6),
            "DAI": ("0x6B175474E89094C44Da98b954EedeAC495271d0F", 18),
        },
    ),
    "arb": EvmChain(
        "arb", "Arbitrum", 42161, "https://arbitrum-one-rpc.publicnode.com", "ETH",
        {
            "USDC": ("0xaf88d065e77c8cC2239327C5EDb3A432268e5831", 6),
            "USDC.e": ("0xFF970A61A04b1cA14834A43f5dE4533eBDDB5CC8", 6),
            "USDT": ("0xFd086bC7CD5C481DCC9C85ebE478A1C0b69FCbb9", 6),
        },
    ),
    "base": EvmChain(
        "base", "Base", 8453, "https://base-rpc.publicnode.com", "ETH",
        {
            "USDC": ("0x833589fCD6eDb6E08f4c7C32D4f71b54bdA02913", 6),
            "USDT": ("0xfde4C96c8593536E31F229EA8f37b2ADa2699bb2", 6),
        },
    ),
    "polygon": EvmChain(
        "polygon", "Polygon", 137, "https://polygon-bor-rpc.publicnode.com", "POL",
        {
            "USDC": ("0x3c499c542cEF5E3811e1192ce70d8cC03d5c3359", 6),
            "USDC.e": ("0x2791Bca1f2de4661ED88A30C99A7a9449Aa84174", 6),
            "USDT": ("0xc2132D05D31c914a87C6611C10748AEb04B58e8F", 6),
            "WETH": ("0x7ceB23fD6bC0adD59E62ac25578270cFf1b9f619", 18),
        },
    ),
    "avax": EvmChain(
        "avax", "Avalanche", 43114, "https://avalanche-c-chain-rpc.publicnode.com", "AVAX",
        {
            "USDC": ("0xB97EF9Ef8734C71904D8002F8b6Bc66Dd9c48a6E", 6),
            "USDT": ("0x9702230A8Ea53601f5cD2dc00fDBc13d4dF4A8c7", 6),
            "WETH.e": ("0x49D5c2BdFfac6CE2BFdB6640F4F80f226bc10bAB", 18),
        },
    ),
    "op": EvmChain(
        "op", "OP Mainnet", 10, "https://optimism-rpc.publicnode.com", "ETH",
        {
            "USDC": ("0x0b2C639c533813f4Aa9D7837CAf62653d097Ff85", 6),
            "USDC.e": ("0x7F5c764cBc14f9669B88837ca1490cCa17c31607", 6),
            "USDT": ("0x94b008aA00579c1307B0EF2c499aD98a8ce58e58", 6),
        },
    ),
}

CHAIN_ALIASES = {
    "ethereum": "eth", "mainnet": "eth",
    "arbitrum": "arb", "arbitrum-one": "arb",
    "matic": "polygon", "pol": "polygon",
    "avalanche": "avax", "avalanche-c": "avax",
    "optimism": "op", "op-mainnet": "op",
    "tron": "trx", "bitcoin": "btc",
}
NON_EVM = ("trx", "btc")

# TRC20 tokens with known decimals — others are printed as-is
TRC20_KNOWN = {
    "TR7NHqjeKQxGTCi8q8ZY4pL8otSzgjLj6t": ("USDT", 6),
    "TEkxiTehnzSmSe2XqrBj4w32RUN966rdz8": ("USDC", 6),
}

BTC_P2PKH_VER = CoinsConf.BitcoinMainNet.ParamByKey("p2pkh_net_ver")
BTC_P2SH_VER = CoinsConf.BitcoinMainNet.ParamByKey("p2sh_net_ver")
BTC_WIF_VER = CoinsConf.BitcoinMainNet.ParamByKey("wif_net_ver")


# ------------------------------------------------------------------- data structures


@dataclass
class Row:
    network: str
    scheme: str
    path: str
    index: int
    address: str
    symbol: str = ""
    balance: Decimal = Decimal(0)
    tx_count: int | None = None
    tokens: str = ""
    used: bool = False
    status: str = ""
    priv: str = ""
    matched: bool = False  # matched via --find
    checker: Callable[["Fetcher", "Row"], None] | None = field(default=None, repr=False)

    @property
    def is_error(self) -> bool:
        return self.status.startswith("ERROR")

    @property
    def status_text(self) -> str:
        """Status for the report: the --find flag and an RPC error are shown together."""
        parts = ["found via --find"] if self.matched else []
        if self.status:
            parts.append(self.status)
        return "; ".join(parts)

    @property
    def has_assets(self) -> bool:
        return self.balance > 0 or bool(self.tokens)

    @property
    def is_empty(self) -> bool:
        """Empty address: no assets, no history, no error, not found via --find."""
        return not (self.has_assets or self.used or self.is_error or self.matched)


# ------------------------------------------------------------------- derivation


def derive_evm(seed: bytes, account: int, count: int) -> list[tuple[str, str, str]]:
    return _bip_addresses(Bip44, Bip44Coins.ETHEREUM, seed, account, count, 44, 60)


def _bip_addresses(cls, coin, seed: bytes, account: int, count: int,
                   purpose: int, coin_type: int) -> list[tuple[str, str, str]]:
    acc = cls.FromSeed(seed, coin).Purpose().Coin().Account(account).Change(
        Bip44Changes.CHAIN_EXT
    )
    out = []
    for i in range(count):
        ctx = acc.AddressIndex(i)
        out.append((f"m/{purpose}'/{coin_type}'/{account}'/0/{i}",
                    ctx.PublicKey().ToAddress(),
                    ctx.PrivateKey().ToWif()))
    return out


def _btc_raw_addresses(seed: bytes, count: int, encoder, net_ver) -> list[tuple[str, str, str]]:
    """BIP32 and BIP141 tabs from iancoleman: path m/0, index is the last level."""
    base = Bip32Slip10Secp256k1.FromSeed(seed).DerivePath("m/0")
    out = []
    for i in range(count):
        node = base.ChildKey(i)
        out.append((f"m/0/{i}",
                    encoder.EncodeKey(node.PublicKey().KeyObject(), net_ver=net_ver),
                    WifEncoder.Encode(node.PrivateKey().KeyObject(), net_ver=BTC_WIF_VER)))
    return out


def generate_rows(seed: bytes, account: int, chains: list[str], count: int,
                  with_keys: bool) -> list[Row]:
    rows: list[Row] = []

    def add(network: str, scheme: str, symbol: str, checker,
            items: list[tuple[str, str, str]]):
        for i, (path, addr, priv) in enumerate(items):
            rows.append(Row(network, scheme, path, i, addr, symbol=symbol,
                            priv=priv if with_keys else "", checker=checker))

    evm_keys = [k for k in EVM_CHAINS if k in chains]
    if evm_keys:
        evm = derive_evm(seed, account, count)
        for key in evm_keys:
            ch = EVM_CHAINS[key]
            add(ch.name, "BIP44", ch.symbol, partial(check_evm, ch), evm)

    if "trx" in chains:
        add("Tron", "BIP44", "TRX", check_trx,
            _bip_addresses(Bip44, Bip44Coins.TRON, seed, account, count, 44, 195))

    if "btc" in chains:
        add("Bitcoin", "BIP44", "BTC", check_btc,
            _bip_addresses(Bip44, Bip44Coins.BITCOIN, seed, account, count, 44, 0))
        add("Bitcoin", "BIP49", "BTC", check_btc,
            _bip_addresses(Bip49, Bip49Coins.BITCOIN, seed, account, count, 49, 0))
        add("Bitcoin", "BIP84", "BTC", check_btc,
            _bip_addresses(Bip84, Bip84Coins.BITCOIN, seed, account, count, 84, 0))
        add("Bitcoin", "BIP32", "BTC", check_btc,
            _btc_raw_addresses(seed, count, P2PKHAddrEncoder, BTC_P2PKH_VER))
        add("Bitcoin", "BIP141", "BTC", check_btc,
            _btc_raw_addresses(seed, count, P2SHAddrEncoder, BTC_P2SH_VER))

    return rows


# ----------------------------------------------------------------------- network


class Fetcher:
    """HTTP client with a global delay between requests."""

    def __init__(self, delay: float):
        self.delay = delay
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": USER_AGENT})
        self._last = 0.0
        self.count = 0

    def _wait(self):
        if self._last:
            rest = self.delay - (time.monotonic() - self._last)
            if rest > 0:
                time.sleep(rest)
        self._last = time.monotonic()

    def request(self, method: str, url: str, **kw):
        last_err = None
        for _ in range(HTTP_RETRIES):
            self._wait()
            self.count += 1
            try:
                r = self.session.request(method, url, timeout=HTTP_TIMEOUT, **kw)
                if r.status_code == 429 or r.status_code >= 500:
                    last_err = f"HTTP {r.status_code}"
                    continue
                r.raise_for_status()
                return r
            except requests.RequestException as e:
                last_err = f"{type(e).__name__}: {e}"
        raise RuntimeError(last_err or "request failed")


def hex_to_int(v) -> int:
    if isinstance(v, str) and v.startswith("0x"):
        return int(v, 16) if len(v) > 2 else 0
    return int(v or 0)


def fmt_amount(raw: int, decimals: int) -> str:
    return f"{Decimal(raw) / Decimal(10 ** decimals):f}"


def check_evm(chain: EvmChain, f: Fetcher, row: Row) -> None:
    """Balance, nonce, and all network tokens — in a single batch request."""
    calls = [
        {"jsonrpc": "2.0", "id": 0, "method": "eth_getBalance",
         "params": [row.address, "latest"]},
        {"jsonrpc": "2.0", "id": 1, "method": "eth_getTransactionCount",
         "params": [row.address, "latest"]},
    ]
    order = []
    data = "0x70a08231" + "0" * 24 + row.address[2:].lower()  # balanceOf(address)
    for name, (contract, dec) in chain.tokens.items():
        rid = len(calls)
        order.append((rid, name, dec))
        calls.append({"jsonrpc": "2.0", "id": rid, "method": "eth_call",
                      "params": [{"to": contract, "data": data}, "latest"]})

    res = f.request("POST", chain.rpc, json=calls).json()
    by_id = {x["id"]: x for x in res} if isinstance(res, list) else {0: res}

    def val(i):
        item = by_id.get(i, {})
        if "error" in item:
            raise RuntimeError(item["error"].get("message", "rpc error"))
        return hex_to_int(item.get("result"))

    row.balance = Decimal(val(0)) / Decimal(10 ** 18)
    row.tx_count = val(1)

    parts = []
    for rid, name, dec in order:
        try:
            amount = val(rid)
        except RuntimeError:
            continue
        if amount:
            parts.append(f"{name}={fmt_amount(amount, dec)}")
    row.tokens = ", ".join(parts)
    row.used = bool(row.balance or row.tx_count or row.tokens)


def check_trx(f: Fetcher, row: Row) -> None:
    r = f.request("GET", f"{TRON_API}/v1/accounts/{row.address}")
    data = (r.json() or {}).get("data") or []
    if not data:
        row.balance = Decimal(0)
        row.status = "not activated"
        return

    acc = data[0]
    row.balance = Decimal(acc.get("balance", 0)) / Decimal(10 ** 6)

    parts = []
    for entry in acc.get("trc20") or []:
        for contract, amount in entry.items():
            amount = int(amount)
            if not amount:
                continue
            name, dec = TRC20_KNOWN.get(contract, (contract[:8] + "…", 6))
            parts.append(f"{name}={fmt_amount(amount, dec)}")
    row.tokens = ", ".join(parts)
    row.used = True


def check_btc(f: Fetcher, row: Row) -> None:
    r = f.request("GET", f"{BTC_API}/address/{row.address}")
    d = r.json()
    chain, mem = d.get("chain_stats", {}), d.get("mempool_stats", {})
    sats = (chain.get("funded_txo_sum", 0) - chain.get("spent_txo_sum", 0)
            + mem.get("funded_txo_sum", 0) - mem.get("spent_txo_sum", 0))
    row.balance = Decimal(sats) / Decimal(10 ** 8)
    row.tx_count = chain.get("tx_count", 0) + mem.get("tx_count", 0)
    row.used = bool(row.tx_count or sats)


# ------------------------------------------------- report name and progress indicator


def plural_rows(n: int) -> str:
    return "row" if n == 1 else "rows"


def fmt_duration(seconds: float) -> str:
    total = int(seconds)
    h, rest = divmod(total, 3600)
    m, s = divmod(rest, 60)
    if h:
        return f"{h}h {m:02d}m"
    if m:
        return f"{m}m {s:02d}s"
    return f"{s}s"


def default_out_name(chains: list[str], count: int) -> str:
    """seed_report_2026-09-23_14-05-33_all_n500.xlsx — reports never overwrite each other."""
    keys = set(chains)
    if keys == set(EVM_CHAINS) | set(NON_EVM):
        tag = "all"
    elif keys == set(EVM_CHAINS):
        tag = "evm"
    else:
        tag = "-".join(chains)
    return f"seed_report_{time.strftime('%Y-%m-%d_%H-%M-%S')}_{tag}_n{count}.xlsx"


def unique_path(path: str, overwrite: bool) -> tuple[str, str | None]:
    """Returns (path, message). An existing file is never overwritten."""
    if overwrite or not os.path.exists(path):
        return path, None
    stem, ext = os.path.splitext(path)
    for i in range(2, 1000):
        candidate = f"{stem}_{i}{ext}"
        if not os.path.exists(candidate):
            return candidate, (f"File {path} already exists, saving as {candidate} instead "
                               f"(use --overwrite to overwrite).")
    raise RuntimeError(f"could not find a free name next to {path}")


class Progress:
    """Live progress line; permanent lines are printed only for notable addresses."""

    def __init__(self, total: int):
        self.total = total
        self.tty = sys.stdout.isatty()
        self.start = time.monotonic()
        self.width = 0

    def _clear(self) -> None:
        if self.tty and self.width:
            sys.stdout.write("\r" + " " * self.width + "\r")
            self.width = 0

    def line(self, text: str, err: bool = False) -> None:
        self._clear()
        print(text, file=sys.stderr if err else sys.stdout, flush=True)

    def tick(self, n: int, found: int) -> None:
        elapsed = time.monotonic() - self.start
        eta = elapsed / n * (self.total - n) if n else 0.0
        text = (f"[{n}/{self.total}] {n / self.total * 100:5.1f}% | "
                f"elapsed {fmt_duration(elapsed)} | remaining ~{fmt_duration(eta)} | "
                f"with assets: {found}")
        if self.tty:
            sys.stdout.write("\r" + text)
            sys.stdout.flush()
            self.width = len(text)
        elif n % 25 == 0 or n == self.total:
            print(text, flush=True)

    def finish(self) -> None:
        self._clear()


# ---------------------------------------------------------------------- output


def write_xlsx(rows: list[Row], path: str, with_keys: bool, all_rows: list[Row]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Addresses"

    headers = ["Network", "Scheme", "Derivation Path", "#", "Address",
               "Balance", "Coin", "Tx Count", "Tokens", "Used", "Status"]
    if with_keys:
        headers.append("Private Key (WIF/hex)")
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="1F3864")
    hit_fill = PatternFill("solid", fgColor="C6EFCE")   # has assets
    used_fill = PatternFill("solid", fgColor="FFF2CC")  # had activity
    err_fill = PatternFill("solid", fgColor="F8CBAD")   # request error
    find_fill = PatternFill("solid", fgColor="BDD7EE")  # found via --find

    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        data = [row.network, row.scheme, row.path, row.index, row.address,
                float(row.balance), row.symbol,
                row.tx_count if row.tx_count is not None else "",
                row.tokens, "yes" if row.used else "no", row.status_text]
        if with_keys:
            data.append(row.priv)
        ws.append(data)

        cells = ws[ws.max_row]
        cells[5].number_format = "0.00000000"
        if row.matched:
            fill = find_fill
        elif row.is_error:
            fill = err_fill
        elif row.has_assets:
            fill = hit_fill
        elif row.used:
            fill = used_fill
        else:
            fill = None
        if fill:
            for c in cells:
                c.fill = fill

    widths = [13, 9, 22, 5, 46, 18, 8, 12, 40, 13, 26, 56]
    for i, w in enumerate(widths[:len(headers)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---- Summary: computed over the FULL set of addresses, not the filtered one
    s = wb.create_sheet("Summary")
    s.append(["Network", "Scheme", "Checked", "With History", "With Assets",
              "Empty (hidden)", "Errors", "Total Balance", "Coin"])
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    groups: dict[tuple[str, str], list[Row]] = {}
    for row in all_rows:
        groups.setdefault((row.network, row.scheme), []).append(row)

    shown = {id(r) for r in rows}
    for (net, scheme), items in groups.items():
        total = sum((r.balance for r in items), Decimal(0))
        s.append([net, scheme, len(items),
                  sum(1 for r in items if r.used),
                  sum(1 for r in items if r.has_assets),
                  sum(1 for r in items if id(r) not in shown),
                  sum(1 for r in items if r.is_error),
                  float(total), items[0].symbol])
        s.cell(s.max_row, 8).number_format = "0.00000000"

    for i, w in enumerate([13, 9, 11, 12, 12, 16, 9, 20, 8], start=1):
        s.column_dimensions[get_column_letter(i)].width = w
    s.freeze_panes = "A2"
    s.auto_filter.ref = s.dimensions

    wb.save(path)


# ----------------------------------------------------------------- --find search


@dataclass
class FindResult:
    target: str             # as entered by the user
    hits: list[Row] = field(default_factory=list)

    @property
    def key(self) -> str:
        return self.target.lower()


def search_addresses(rows: list[Row], spec: str) -> list[FindResult]:
    """Flags rows that matched --find. Runs over the FULL set of addresses."""
    results: list[FindResult] = []
    seen: set[str] = set()
    for raw in spec.split(","):
        target = raw.strip()
        if target and target.lower() not in seen:
            seen.add(target.lower())
            results.append(FindResult(target))

    by_key = {r.key: r for r in results}
    for row in rows:
        res = by_key.get(row.address.lower())
        if res is not None:
            row.matched = True
            res.hits.append(row)
    return results


def print_find_report(results: list[FindResult], rows: list[Row],
                      checked: bool, args) -> None:
    """Search result notice: printed both after generation and at the end."""
    line = "─" * 78
    n_found = sum(1 for r in results if r.hits)
    print(f"\n{line}")
    print(f"SEARCH --find: found {n_found} of {len(results)} "
          f"among {len(rows)} generated addresses")
    print(line)

    for res in results:
        if not res.hits:
            print(f"  ✗ NOT FOUND  {res.target}")
            continue

        # a single EVM address matches multiple networks at once — group by path
        groups: dict[tuple[str, str, int, str], list[Row]] = {}
        for r in res.hits:
            groups.setdefault((r.scheme, r.path, r.index, r.address), []).append(r)

        for (scheme, path, index, address), items in groups.items():
            print(f"  ✓ FOUND      {address}")
            print(f"                {scheme}  {path}  (index {index})")
            print(f"                networks: {', '.join(i.network for i in items)}")
            if checked:
                assets = [f"{i.network} {i.balance:.8f} {i.symbol}"
                          + (f", {i.tokens}" if i.tokens else "")
                          for i in items if i.has_assets]
                errs = [i.network for i in items if i.is_error]
                print(f"                assets: {'; '.join(assets) if assets else 'none'}")
                if errs:
                    print(f"                not checked (RPC error): {', '.join(errs)}")

    if n_found < len(results):
        schemes = len({(r.network, r.scheme) for r in rows})
        pw = "set" if args.passphrase else "not set"
        print(f"  Checked {schemes} scheme/network combos × {args.count} addresses, "
              f"account' = {args.account}, passphrase {pw}.")
        print(f"  If the address does belong to this phrase — increase -n, "
              f"check -a and -p, and the network set in -c.")
    print(line)


# ----------------------------------------------------------------------- main


def read_mnemonic(args) -> str:
    if args.mnemonic:
        return " ".join(args.mnemonic.split())
    if args.mnemonic_file:
        with open(args.mnemonic_file, encoding="utf-8") as fh:
            return " ".join(fh.read().split())
    import getpass
    return " ".join(getpass.getpass("Seed phrase (hidden input): ").split())


def parse_chains(spec: str) -> tuple[list[str], list[str]]:
    """Returns (selected keys in canonical order, unknown ones)."""
    order = list(EVM_CHAINS) + list(NON_EVM)
    wanted: set[str] = set()
    unknown: list[str] = []
    for raw in spec.split(","):
        key = CHAIN_ALIASES.get(raw.strip().lower(), raw.strip().lower())
        if not key:
            continue
        if key == "evm":
            wanted.update(EVM_CHAINS)
        elif key == "all":
            wanted.update(order)
        elif key in order:
            wanted.add(key)
        else:
            unknown.append(raw.strip())
    return [k for k in order if k in wanted], unknown


def main() -> int:
    p = argparse.ArgumentParser(
        description="Check the first N addresses derived from a BIP39 seed phrase: "
                    "EVM networks (Ethereum, Arbitrum, Base, Polygon, Avalanche, OP), Tron, Bitcoin.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Networks for -c: " + ", ".join(list(EVM_CHAINS) + list(NON_EVM))
               + "; groups: evm, all",
    )
    p.add_argument("-m", "--mnemonic", help="seed phrase, 12-24 words (in quotes)")
    p.add_argument("-f", "--mnemonic-file", help="file containing the seed phrase")
    p.add_argument("-p", "--passphrase", default="", help="BIP39 passphrase (extra word)")
    p.add_argument("-n", "--count", type=int, default=50,
                   help="how many addresses per scheme (default 50)")
    p.add_argument("-d", "--delay", type=float, default=3.0,
                   help="delay between requests, in seconds (default 3)")
    p.add_argument("-a", "--account", type=int, default=0,
                   help="account' number for BIP44/49/84 (default 0)")
    p.add_argument("-c", "--chains", default="all",
                   help="comma-separated networks or the evm/all group (default all)")
    p.add_argument("-o", "--out",
                   help="xlsx file name; auto-generated by default as "
                        "seed_report_<date>_<time>_<networks>_n<count>.xlsx")
    p.add_argument("--overwrite", action="store_true",
                   help="overwrite the existing file instead of adding a suffix")
    p.add_argument("-v", "--verbose", action="store_true",
                   help="print every checked address, not just the notable ones")
    p.add_argument("--find", metavar="ADDR",
                   help="search for address(es), comma-separated, among the generated ones; "
                        "search runs over the full set, before hiding empty rows")
    p.add_argument("--keep-empty", action="store_true",
                   help="don't hide rows with zero assets and no history in the report")
    p.add_argument("--with-keys", action="store_true",
                   help="WARNING: export private keys to the xlsx")
    p.add_argument("--dry-run", action="store_true",
                   help="only generate addresses, without contacting the network")
    p.add_argument("--no-validate", action="store_true",
                   help="skip the BIP39 checksum validation")
    args = p.parse_args()

    mnemonic = read_mnemonic(args)
    words = mnemonic.split()
    if len(words) not in (12, 15, 18, 21, 24):
        print(f"Error: expected 12/15/18/21/24 words, got {len(words)}.", file=sys.stderr)
        return 2
    if not args.no_validate:
        try:
            Bip39MnemonicValidator().Validate(mnemonic)
        except Exception as e:
            print(f"Error: invalid seed phrase ({e}). Disable the check with: --no-validate",
                  file=sys.stderr)
            return 2

    chains, unknown = parse_chains(args.chains)
    if unknown:
        print(f"Error: unknown networks: {', '.join(unknown)}", file=sys.stderr)
        return 2
    if not chains:
        print("Error: no networks selected.", file=sys.stderr)
        return 2

    if not args.out:
        args.out = default_out_name(chains, args.count)

    seed = Bip39SeedGenerator(mnemonic).Generate(args.passphrase)
    all_rows = generate_rows(seed, args.account, chains, args.count, args.with_keys)

    print(f"Words: {len(words)} | addresses: {len(all_rows)}")
    seen = set()
    for r in all_rows:
        if (r.network, r.scheme) not in seen:
            seen.add((r.network, r.scheme))
            print(f"  {r.network:11s} {r.scheme:6s} {r.path}")

    # ---- search requested addresses over the FULL set, before hiding empty rows
    find_results: list[FindResult] = []
    if args.find:
        find_results = search_addresses(all_rows, args.find)
        print_find_report(find_results, all_rows, checked=False, args=args)

    if not args.dry_run:
        eta = len(all_rows) * args.delay
        print(f"Delay {args.delay:g}s between requests, "
              f"estimated time ≈ {eta / 60:.1f} min.\n")
        f = Fetcher(args.delay)
        prog = Progress(len(all_rows))
        found = 0
        for n, row in enumerate(all_rows, start=1):
            try:
                row.checker(f, row)
                if row.has_assets:
                    found += 1
                if args.verbose or row.used or row.matched:
                    mark = "•" if row.used else " "
                    extra = f"  {row.tokens}" if row.tokens else ""
                    prog.line(f"  {mark} {row.network:11s} {row.scheme:6s} "
                              f"{row.address:44s} {row.balance:.8f} {row.symbol}{extra}")
            except Exception as e:
                row.status = f"ERROR: {e}"
                prog.line(f"  ! {row.network:11s} {row.scheme:6s} "
                          f"{row.address:44s} {row.status}", err=True)
            prog.tick(n, found)
        prog.finish()
    else:
        print("(dry-run: network is not queried)\n")
        if args.find:
            print(f"Full list of {len(all_rows)} addresses is in the report.")
        else:
            for row in all_rows:
                print(f"{row.network:11s} {row.scheme:6s} {row.path:20s} {row.address}")

    # ---- hide empty rows (nothing to filter in dry-run)
    if args.keep_empty or args.dry_run:
        rows = all_rows
    else:
        rows = [r for r in all_rows if not r.is_empty]

    out_path, renamed = unique_path(args.out, args.overwrite)
    if renamed:
        print(f"\n{renamed}")
    print(f"\nBuilding report ({len(rows)} {plural_rows(len(rows))})…")
    write_xlsx(rows, out_path, args.with_keys, all_rows)

    hits = [r for r in all_rows if r.has_assets]
    used = [r for r in all_rows if r.used]
    errs = [r for r in all_rows if r.is_error]
    hidden = len(all_rows) - len(rows)
    print(f"\nDone. Addresses with history: {len(used)}, with assets: {len(hits)}, "
          f"errors: {len(errs)}, empty hidden: {hidden}")
    for r in hits:
        print(f"  {r.network} {r.scheme} {r.path}  {r.address}  "
              f"{r.balance:.8f} {r.symbol} {r.tokens}")
    if find_results:
        print_find_report(find_results, all_rows, checked=not args.dry_run, args=args)
    print(f"\nReport: {os.path.abspath(out_path)} "
          f"({len(rows)} {plural_rows(len(rows))})")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nInterrupted by user.", file=sys.stderr)
        sys.exit(130)
